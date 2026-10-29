"""Closed-loop synthetic-history simulation for the training advisor.

Simulates a trainee who always follows advise()'s top-ranked pick, three
sessions/week, one heavy piece then four workout pieces per session -
docs/GUIDE.md's "Capacity arithmetic" section works out that shape
from the stated goals, though it's presented there as guidance for the
human's own session design, not something advisor.py enforces. Simulation
needs *a* concrete shape to run at all, so it mirrors that math.

This is a self-consistency check, not real data: does following the
model's own advice for two months produce balanced coverage, or does some
axis starve/dominate? It's the "comparison harness" GUIDE.md's parameter
notes gesture at, run before any real logging exists. It cannot validate
the exercise dictionary's E/R weights or the tuned constants against
reality - only real logs can do that (see docs/AUDIT.md).

The database is entirely in-memory and synthetic; nothing personal is
written anywhere. Only the report is committed.

Usage:
    python -m training.simulate
"""

from __future__ import annotations

import math
import random
from collections import Counter
from dataclasses import dataclass, field
from datetime import date as Date, timedelta
from pathlib import Path

from . import advisor, db

REPO_ROOT = Path(__file__).resolve().parents[2]
REPORT_PATH = REPO_ROOT / "data" / "simulation_report.md"
SCHEDULE_XLSX_PATH = REPO_ROOT / "data" / "simulation_schedule.xlsx"

SESSION_WEEKDAYS = (0, 2, 4)  # Mon/Wed/Fri
SIMULATION_DAYS = 60
SEED = 20260101

REP_RANGE_WORKOUT = (6, 15)
REP_RANGE_HEAVY_FALLBACK = (3, 8)  # used only until suggest_rep_count() has history


@dataclass
class SessionResult:
    date: Date
    picks: list[advisor.Candidate]


@dataclass
class SimulationResult:
    dictionary: advisor.Dictionary
    conn: object
    sessions: list[SessionResult]
    weekly_verdicts: list[tuple[Date, advisor.Verdict]]
    start: Date
    end: Date
    user_id: int = 0
    empty_advise_events: list[tuple[Date, str]] = field(default_factory=list)


def session_dates(start: Date, days: int, weekdays: tuple[int, ...] = SESSION_WEEKDAYS) -> list[Date]:
    return [
        start + timedelta(days=i)
        for i in range(days)
        if (start + timedelta(days=i)).weekday() in weekdays
    ]


def _log_one_or_two_sides(conn, block_id, ex, dict_version, reps, scheme, rng) -> None:
    if ex.laterality == "unilateral":
        db.log_set(conn, block_id, ex.id, dict_version, reps=reps, scheme=scheme, side="L")
        db.log_set(conn, block_id, ex.id, dict_version, reps=reps, scheme=scheme, side="R")
    else:
        db.log_set(conn, block_id, ex.id, dict_version, reps=reps, scheme=scheme)


def _commit_pick(conn, session_id, dictionary, sets_so_far, day, pick, rng) -> None:
    ex = pick.exercise
    block_id = db.log_block(conn, session_id, pick.style)
    if pick.style == "heavy":
        scheme = advisor.suggest_rep_scheme(dictionary, sets_so_far, day, pick.pattern)
        reps = advisor.suggest_rep_count(sets_so_far, ex.id, "heavy") or rng.randint(*REP_RANGE_HEAVY_FALLBACK)
        n_sets = 5 if "5" in scheme else 3
        for _ in range(n_sets):
            _log_one_or_two_sides(conn, block_id, ex, dictionary.version, reps, scheme, rng)
    else:
        reps = rng.randint(*REP_RANGE_WORKOUT)
        n_sets = rng.randint(1, 3)
        for _ in range(n_sets):
            _log_one_or_two_sides(conn, block_id, ex, dictionary.version, reps, None, rng)


def run_session(dictionary, conn, user_id, day, rng, empty_advise_events) -> SessionResult:
    chosen: list[str] = []
    picks: list[advisor.Candidate] = []
    session_id = db.log_session(conn, user_id, day.isoformat())

    sets_so_far = db.fetch_sets(conn, user_id)
    heavy_candidates = advisor.advise(dictionary, sets_so_far, day, "heavy", chosen=chosen)
    if not heavy_candidates:
        empty_advise_events.append((day, "heavy"))
    else:
        pick = heavy_candidates[0]
        _commit_pick(conn, session_id, dictionary, sets_so_far, day, pick, rng)
        chosen.append(pick.exercise.id)
        picks.append(pick)

    for _ in range(4):
        sets_so_far = db.fetch_sets(conn, user_id)  # re-fetch: this session's own picks must be visible
        workout_candidates = advisor.advise(dictionary, sets_so_far, day, "workout", chosen=chosen)
        if not workout_candidates:
            empty_advise_events.append((day, "workout"))
            break
        pick = workout_candidates[0]
        _commit_pick(conn, session_id, dictionary, sets_so_far, day, pick, rng)
        chosen.append(pick.exercise.id)
        picks.append(pick)

    return SessionResult(date=day, picks=picks)


def simulate(days: int = SIMULATION_DAYS, start: Date | None = None, seed: int = SEED) -> SimulationResult:
    rng = random.Random(seed)
    dictionary = advisor.load_dictionary()
    conn = db.connect()
    # Synthetic single user - schema.sql requires every session to belong to
    # one (see docs/STATUS.md's webapp section). Not a real login;
    # just satisfies the same foreign key a real webapp user would.
    user_id = db.create_user(conn, "simulated@example.com", "not-a-real-hash")
    start = start or (Date.today() - timedelta(days=days))

    dates = session_dates(start, days)
    sessions = []
    weekly_verdicts = []
    empty_advise_events: list[tuple[Date, str]] = []

    for day in dates:
        result = run_session(dictionary, conn, user_id, day, rng, empty_advise_events)
        sessions.append(result)
        if day.weekday() == 0:
            verdict = advisor.recommend(dictionary, db.fetch_sets(conn, user_id), day)
            weekly_verdicts.append((day, verdict))

    return SimulationResult(
        dictionary=dictionary, conn=conn, sessions=sessions, weekly_verdicts=weekly_verdicts,
        start=start, end=dates[-1] if dates else start, empty_advise_events=empty_advise_events,
        user_id=user_id,
    )


# --- Report ------------------------------------------------------------------

def _format_report(result: SimulationResult) -> str:
    dictionary = result.dictionary
    all_sets = db.fetch_sets(result.conn, result.user_id)
    final_stats = advisor.build_log_stats(dictionary, all_sets, result.end)
    final_scores = {
        cell: advisor.ewma_score(ewma) for cell, ewma in final_stats.cell_ewma.items()
    }

    dated_picks = [(s.date, p) for s in result.sessions for p in s.picks]
    all_picks = [p for _d, p in dated_picks]

    pick_counts = Counter(p.exercise.id for p in all_picks)
    n_admissible_exercises = len({
        ex.id for ex in dictionary.exercises.values()
        if any(ex.E.get(p, 0.0) > 0 for p in dictionary.patterns)
    })

    lines = []
    lines.append("# Training advisor - two-month synthetic simulation report\n")
    lines.append(
        "Not real data. A closed-loop check: the simulated trainee always follows "
        "`advise()`'s top pick, three sessions/week, one heavy piece then four "
        "workout pieces (see `docs/GUIDE.md`, \"Capacity arithmetic\"). "
        "Regenerate with `python -m training.simulate`. See `docs/AUDIT.md` "
        "for which constants this can and can't validate.\n"
    )
    lines.append(f"- Dictionary version: {dictionary.version}")
    lines.append(f"- Period: {result.start.isoformat()} to {result.end.isoformat()} ({len(result.sessions)} sessions)")
    lines.append(f"- Total pieces logged: {len(all_picks)}")
    lines.append(f"- Distinct exercises used: {len(pick_counts)} / {n_admissible_exercises} pattern-bearing exercises in the dictionary")
    if result.empty_advise_events:
        lines.append(f"- **{len(result.empty_advise_events)} session(s) had no candidates for a requested mode** - see below")
    lines.append("")

    lines.append("## Final EWMA scores per (pattern, style) cell\n")
    lines.append(
        "score = -log(ewma / MEMORY_K). 0 = exactly on target, positive = "
        "behind (grows without bound as neglect continues; `inf` means never "
        "trained at that style at all), negative = ahead. Sorted worst "
        "(most behind) first.\n"
    )
    lines.append("| pattern | style | score |")
    lines.append("|---|---|---:|")
    for (pattern, style), value in sorted(final_scores.items(), key=lambda kv: -kv[1]):
        lines.append(f"| {pattern} | {style} | {value:+.2f} |" if math.isfinite(value) else f"| {pattern} | {style} | inf |")
    lines.append("")

    lines.append("## Longest gaps at end of simulation\n")
    lines.append(
        "No hard overdue clock exists any more - a neglected cell's score "
        "grows without bound on its own (see `docs/GUIDE.md`'s "
        "Scoring section), but nothing forces a specific cell above the "
        "others by a fixed date. This table is the sanity check for that "
        "removal: days since each admissible cell was literally last "
        "trained, as of the end of the run. Sorted worst first - if "
        "anything here looks uncomfortably large, that's the signal the "
        "clock used to suppress by construction rather than by ranking.\n"
    )
    gaps = []
    for pattern, style in advisor.admissible_cells(dictionary):
        last = final_stats.cell_last_seen.get((pattern, style))
        gap = (result.end - last).days if last else math.inf
        gaps.append((pattern, style, gap))
    gaps.sort(key=lambda row: -row[2])

    lines.append("| pattern | style | days since last touched |")
    lines.append("|---|---|---:|")
    for pattern, style, gap in gaps:
        lines.append(f"| {pattern} | {style} | {'never' if math.isinf(gap) else gap} |")
    lines.append("")

    lines.append("## Movement variety\n")
    lines.append("Top 10 most-picked exercises:\n")
    lines.append("| exercise | times picked |")
    lines.append("|---|---:|")
    for ex_id, count in pick_counts.most_common(10):
        lines.append(f"| {dictionary.exercises[ex_id].name} | {count} |")
    lines.append("")

    lines.append("## Weekly recommend() verdicts\n")
    lines.append("| Monday | style | band | margin |")
    lines.append("|---|---|---|---:|")
    for day, verdict in result.weekly_verdicts:
        lines.append(f"| {day.isoformat()} | {verdict.style or '-'} | {verdict.band} | {verdict.margin:.1f} |")
    lines.append("")

    if result.empty_advise_events:
        lines.append("## Empty-candidate events\n")
        lines.append("| date | mode |")
        lines.append("|---|---|")
        for day, mode in result.empty_advise_events:
            lines.append(f"| {day.isoformat()} | {mode} |")
        lines.append("")

    return "\n".join(lines) + "\n"


def _pick_session_date(result: SimulationResult, pick: advisor.Candidate) -> Date:
    for s in result.sessions:
        if pick in s.picks:
            return s.date
    raise ValueError("pick not found in any session")  # unreachable given how picks are built


# --- Excel schedule (weeks x weekdays, for eyeballing) ------------------------

_WEEKDAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def _week_start(d: Date) -> Date:
    return d - timedelta(days=d.weekday())


def _format_cell(session: SessionResult | None) -> str:
    if session is None:
        return ""
    return "\n".join(f"[{p.style}] {p.exercise.name}" for p in session.picks)


def export_schedule_xlsx(result: SimulationResult, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    sessions_by_date = {s.date: s for s in result.sessions}

    wb = Workbook()

    info = wb.active
    info.title = "Legend"
    info["A1"] = "Training advisor - simulated schedule"
    info["A1"].font = Font(bold=True, size=14)
    info["A3"] = "Not real data - a closed-loop simulation where the trainee always follows advise()'s top pick."
    info["A4"] = f"Dictionary version: {result.dictionary.version}"
    info["A5"] = f"Period: {result.start.isoformat()} to {result.end.isoformat()}"
    info["A6"] = f"Sessions: {len(result.sessions)}   Pieces logged: {sum(len(s.picks) for s in result.sessions)}"
    info["A7"] = "Cell format: one movement per line, tagged [heavy] / [light] / [hhr]."
    info["A8"] = "See data/simulation_report.md for the aggregate numbers behind this schedule."
    for row in range(3, 9):
        info[f"A{row}"].alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 100

    ws = wb.create_sheet("Schedule")
    ws.append(["Week of"] + _WEEKDAY_NAMES)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    week = _week_start(result.start)
    last_week = _week_start(result.end)
    while week <= last_week:
        row = [week.isoformat()]
        for offset in range(7):
            day = week + timedelta(days=offset)
            row.append(_format_cell(sessions_by_date.get(day)))
        ws.append(row)
        week += timedelta(days=7)

    for r in range(2, ws.max_row + 1):
        for c in range(1, 9):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 12
    for col_letter in "BCDEFGH":
        ws.column_dimensions[col_letter].width = 26

    wb.save(path)


def main() -> None:
    result = simulate()

    report = _format_report(result)
    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text(report)
    print(f"wrote {REPORT_PATH}")

    export_schedule_xlsx(result, SCHEDULE_XLSX_PATH)
    print(f"wrote {SCHEDULE_XLSX_PATH}")

    print(f"{len(result.sessions)} sessions, {sum(len(s.picks) for s in result.sessions)} pieces logged")
    if result.empty_advise_events:
        print(f"WARNING: {len(result.empty_advise_events)} empty-candidate events - see report")


if __name__ == "__main__":
    main()
